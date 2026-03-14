import pytest
from pathlib import Path
from docx import Document
from docx.shared import Inches
import fitz
from PIL import Image

from app.services.docx_ad_cleaner import clean_trailing_ad_from_docx
from app.utils.pdf_processor import detect_and_remove_ad, detect_large_ad_image

@pytest.fixture
def temp_dir(tmp_path):
    return tmp_path

@pytest.fixture
def dummy_image(temp_dir):
    image_path = temp_dir / "dummy_ad.jpg"
    img = Image.new('RGB', (800, 600), color=(255, 0, 0))
    img.save(image_path)
    return image_path

def create_docx_with_ad(path, image_path):
    doc = Document()
    doc.add_paragraph("This is normal text paragraph 1.")
    doc.add_paragraph("This is normal text paragraph 2.")
    doc.add_paragraph() # empty paragraph
    p = doc.add_paragraph()
    r = p.add_run()
    r.add_picture(str(image_path), width=Inches(5))
    doc.save(path)

def create_docx_without_ad(path):
    doc = Document()
    doc.add_paragraph("This is normal text paragraph 1.")
    doc.add_paragraph("This is normal text paragraph 2.")
    doc.save(path)

def create_pdf_with_ad(path, image_path):
    doc = fitz.open()
    page = doc.new_page(width=595, height=842) # A4
    page.insert_text((50, 50), "This is normal text.")

    # insert ad image at the bottom
    # bottom 55% starts at 842 * 0.45 = 378.9
    # image width 400 (which is > 40% of 595 = 238)
    # image height 200 (which is > 15% of 842 = 126)
    rect = fitz.Rect(100, 500, 500, 700)
    page.insert_image(rect, filename=str(image_path))

    doc.save(path)
    doc.close()

def create_pdf_without_ad(path):
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((50, 50), "This is normal text.")
    doc.save(path)
    doc.close()

def test_clean_trailing_ad_from_docx_with_ad(temp_dir, dummy_image):
    input_docx = temp_dir / "with_ad.docx"
    output_docx = temp_dir / "cleaned_ad.docx"
    create_docx_with_ad(input_docx, dummy_image)

    # Check that ad exists
    doc = Document(str(input_docx))
    assert len(doc.paragraphs) == 4

    result = clean_trailing_ad_from_docx(input_docx, output_docx)
    assert result is True

    cleaned_doc = Document(str(output_docx))
    # It should have removed the empty paragraph and the image paragraph
    assert len(cleaned_doc.paragraphs) == 2

def test_clean_trailing_ad_from_docx_without_ad(temp_dir):
    input_docx = temp_dir / "no_ad.docx"
    output_docx = temp_dir / "cleaned_no_ad.docx"
    create_docx_without_ad(input_docx)

    result = clean_trailing_ad_from_docx(input_docx, output_docx)
    assert result is False

def test_detect_and_remove_ad_pdf_with_ad(temp_dir, dummy_image):
    input_pdf = temp_dir / "with_ad.pdf"
    create_pdf_with_ad(input_pdf, dummy_image)

    # It should detect and crop the PDF
    result = detect_and_remove_ad(str(input_pdf))
    assert result is True

    doc = fitz.open(str(input_pdf))
    page = doc[0]
    # The cropbox should be smaller than original
    assert page.cropbox.y1 < 842
    doc.close()

def test_detect_and_remove_ad_pdf_without_ad(temp_dir):
    input_pdf = temp_dir / "no_ad.pdf"
    create_pdf_without_ad(input_pdf)

    result = detect_and_remove_ad(str(input_pdf))
    assert result is False
