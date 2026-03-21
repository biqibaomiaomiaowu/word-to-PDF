import argparse
import json
import os
import re
import tempfile
import traceback
from html import unescape
from io import BytesIO
from pathlib import Path


WORKSPACE_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "../../../"))
os.environ["PADDLE_HOME"] = os.path.join(WORKSPACE_ROOT, ".paddle_models")
os.environ["PADDLE_PDX_CACHE_HOME"] = os.path.join(WORKSPACE_ROOT, ".paddlex")
os.environ["PADDLE_INFERENCE_MODEL_DIR"] = os.path.join(WORKSPACE_ROOT, ".paddle_inference")
_HOME_ROOT = os.path.join(WORKSPACE_ROOT, ".cache", "home")
os.makedirs(_HOME_ROOT, exist_ok=True)
os.environ["USERPROFILE"] = _HOME_ROOT
os.environ["HOME"] = _HOME_ROOT
if os.name == "nt":
    _APPDATA = os.path.join(_HOME_ROOT, "AppData", "Roaming")
    _LOCALAPPDATA = os.path.join(_HOME_ROOT, "AppData", "Local")
    os.makedirs(_APPDATA, exist_ok=True)
    os.makedirs(_LOCALAPPDATA, exist_ok=True)
    os.environ["APPDATA"] = _APPDATA
    os.environ["LOCALAPPDATA"] = _LOCALAPPDATA
os.environ.setdefault("PADDLE_PDX_DISABLE_MODEL_SOURCE_CHECK", "True")
os.environ.setdefault("PYTHONUTF8", "1")
os.environ.setdefault("PYTHONIOENCODING", "utf-8")


def _pick_temp_dir() -> str:
    candidates = []
    if os.name == "nt":
        candidates.append(r"C:\Windows\Temp")
    candidates.extend(
        [
            os.path.join(WORKSPACE_ROOT, ".paddlex_runtime_tmp"),
            os.environ.get("TEMP"),
            os.environ.get("TMP"),
        ]
    )

    for candidate in candidates:
        if not candidate:
            continue
        try:
            path = Path(candidate)
            path.mkdir(parents=True, exist_ok=True)
            fd, test_path = tempfile.mkstemp(dir=str(path))
            os.close(fd)
            os.unlink(test_path)
            return str(path)
        except Exception:
            continue

    return tempfile.gettempdir()


_TEMP_DIR = _pick_temp_dir()
os.environ["TEMP"] = _TEMP_DIR
os.environ["TMP"] = _TEMP_DIR


def _has_cached_model(model_name: str) -> bool:
    model_dir = Path(os.environ["PADDLE_PDX_CACHE_HOME"]) / "official_models" / model_name
    return model_dir.is_dir() and any((model_dir / marker).exists() for marker in ("inference.json", "inference.yml"))


def _parse_html_table(html: str) -> list[list[str]]:
    if not html:
        return []

    try:
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(html, "html.parser")
        rows = []
        for tr in soup.find_all("tr"):
            row = [cell.get_text(" ", strip=True) for cell in tr.find_all(["td", "th"])]
            if row:
                rows.append(row)
        return rows
    except Exception:
        return []


def _normalize_space(text: str) -> str:
    return re.sub(r"\s+", " ", (text or "").replace("\xa0", " ")).strip()


def _project_compact_prefix(original: str, compact_prefix: str) -> str:
    if not original or not compact_prefix:
        return ""

    output: list[str] = []
    compact_index = 0
    previous_was_space = False
    for ch in original:
        if ch.isspace():
            if output and not previous_was_space and compact_index < len(compact_prefix):
                output.append(" ")
                previous_was_space = True
            continue

        if compact_index >= len(compact_prefix):
            break

        expected = compact_prefix[compact_index]
        if ch != expected:
            continue

        output.append(ch)
        compact_index += 1
        previous_was_space = False
        if compact_index == len(compact_prefix):
            break

    return "".join(output).strip()


def _collapse_full_repetition(text: str) -> str:
    compact = _normalize_space(text)
    if not compact:
        return ""

    max_unit = len(compact) // 2
    for unit_len in range(1, max_unit + 1):
        if len(compact) % unit_len != 0:
            continue
        unit = compact[:unit_len]
        repeat_count = len(compact) // unit_len
        if repeat_count >= 2 and unit * repeat_count == compact:
            return unit.strip()

    squashed = re.sub(r"\s+", "", compact)
    max_unit = len(squashed) // 2
    for unit_len in range(2, max_unit + 1):
        if len(squashed) % unit_len != 0:
            continue
        unit = squashed[:unit_len]
        repeat_count = len(squashed) // unit_len
        if repeat_count >= 2 and unit * repeat_count == squashed:
            projected = _project_compact_prefix(compact, unit)
            return projected or unit.strip()
    return compact


def _normalize_compare_text(text: str) -> str:
    normalized = _normalize_space(text)
    normalized = re.sub(r"[\s\u3000]+", "", normalized)
    return normalized.casefold()


def _normalize_signature_text(text: str) -> str:
    normalized = _normalize_compare_text(text)
    normalized = re.sub(r"[^\w\u4e00-\u9fff]+", "", normalized)
    return normalized


def _clean_text_content(text: str) -> str:
    text = unescape(text or "")
    if not text:
        return ""

    cleaned_lines: list[str] = []
    for raw_line in re.split(r"[\r\n]+", text):
        line = _collapse_full_repetition(raw_line)
        if not line:
            continue
        if cleaned_lines and _normalize_compare_text(cleaned_lines[-1]) == _normalize_compare_text(line):
            continue
        cleaned_lines.append(line)

    merged = "\n".join(cleaned_lines)
    return _collapse_full_repetition(merged)


def _bbox_overlap_ratio(box_a: list[int], box_b: list[int]) -> float:
    if not box_a or not box_b:
        return 0.0

    ax0, ay0, ax1, ay1 = box_a
    bx0, by0, bx1, by1 = box_b
    inter_x0 = max(ax0, bx0)
    inter_y0 = max(ay0, by0)
    inter_x1 = min(ax1, bx1)
    inter_y1 = min(ay1, by1)
    inter_w = max(0, inter_x1 - inter_x0)
    inter_h = max(0, inter_y1 - inter_y0)
    if inter_w == 0 or inter_h == 0:
        return 0.0

    inter_area = inter_w * inter_h
    area_a = max(1, (ax1 - ax0) * (ay1 - ay0))
    area_b = max(1, (bx1 - bx0) * (by1 - by0))
    return inter_area / float(min(area_a, area_b))


def _bbox_center(box: list[int]) -> tuple[float, float]:
    if not box:
        return 0.0, 0.0
    return (box[0] + box[2]) / 2.0, (box[1] + box[3]) / 2.0


def _deduplicate_page_blocks(blocks: list[dict]) -> list[dict]:
    if not blocks:
        return []

    kept: list[dict] = []
    for block in blocks:
        label = block.get("type")
        norm = block.get("normalized_content", "")
        signature = block.get("signature_content", "") or norm

        if label in {"image", "chart", "seal", "table", "formula"}:
            kept.append(block)
            continue

        if not signature:
            kept.append(block)
            continue

        is_duplicate = False
        for previous in reversed(kept[-16:]):
            prev_signature = previous.get("signature_content", "") or previous.get("normalized_content", "")
            if prev_signature != signature:
                continue

            prev_box = previous.get("bbox") or [0, 0, 0, 0]
            curr_box = block.get("bbox") or [0, 0, 0, 0]
            overlap = _bbox_overlap_ratio(prev_box, curr_box)
            _, prev_cy = _bbox_center(prev_box)
            _, curr_cy = _bbox_center(curr_box)
            same_band = abs(prev_box[1] - curr_box[1]) <= 24
            nearby = abs(prev_cy - curr_cy) <= 64
            if overlap >= 0.6 or same_band or (len(signature) >= 8 and nearby):
                is_duplicate = True
                break

        if not is_duplicate:
            kept.append(block)

    return kept


def _find_mml2omml_xsl() -> Path | None:
    candidates = [
        os.environ.get("MML2OMML_XSL_PATH"),
        r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL",
        r"C:\Program Files (x86)\Microsoft Office\root\Office16\MML2OMML.XSL",
    ]
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            return Path(candidate)
    return None


def _resolve_device(requested_device: str | None) -> str:
    requested = (requested_device or os.environ.get("PADDLE_STRUCTURE_DEVICE") or "auto").strip().lower()
    if requested in {"cpu", "gpu:0"}:
        return requested

    try:
        import paddle

        if paddle.device.is_compiled_with_cuda() and paddle.device.get_device() != "cpu":
            return "gpu:0"
    except Exception:
        pass

    return "cpu"


def _style_map():
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    return {
        "doc_title": {"level": 0, "size": 20, "bold": True, "align": WD_ALIGN_PARAGRAPH.CENTER},
        "header": {"size": 16, "bold": True, "align": WD_ALIGN_PARAGRAPH.CENTER},
        "abstract_title": {"level": 1, "size": 14, "bold": True, "align": WD_ALIGN_PARAGRAPH.CENTER},
        "content_title": {"level": 1, "size": 14, "bold": True, "align": WD_ALIGN_PARAGRAPH.LEFT},
        "reference_title": {"level": 1, "size": 14, "bold": True, "align": WD_ALIGN_PARAGRAPH.LEFT},
        "paragraph_title": {"level": 2, "size": 14, "bold": True, "align": WD_ALIGN_PARAGRAPH.LEFT},
        "abstract": {"size": 12, "align": WD_ALIGN_PARAGRAPH.JUSTIFY},
        "text": {"size": 12, "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "indent": True},
        "content": {"size": 12, "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "indent": True},
        "figure_title": {"size": 10, "align": WD_ALIGN_PARAGRAPH.CENTER},
        "table_title": {"size": 10, "align": WD_ALIGN_PARAGRAPH.CENTER},
        "chart_title": {"size": 10, "align": WD_ALIGN_PARAGRAPH.CENTER},
        "reference": {"size": 12, "align": WD_ALIGN_PARAGRAPH.JUSTIFY},
        "algorithm": {"font": "Courier New", "size": 11, "align": WD_ALIGN_PARAGRAPH.LEFT},
        "formula": {"size": 12, "align": WD_ALIGN_PARAGRAPH.CENTER},
        "vision_footnote": {"size": 9, "align": WD_ALIGN_PARAGRAPH.LEFT},
        "number": {"size": 9, "align": WD_ALIGN_PARAGRAPH.CENTER},
        "footer": {"size": 9, "align": WD_ALIGN_PARAGRAPH.CENTER},
    }


def _normalize_bbox(bbox) -> list[int]:
    if bbox is None:
        return [0, 0, 0, 0]
    return [int(value) for value in bbox]


def _extract_page_blocks(page_result) -> list[dict]:
    blocks = []
    styles = _style_map()
    page_index = int(page_result.get("page_index") or 0)

    for order, block in enumerate(page_result.get("parsing_res_list") or []):
        label = getattr(block, "label", "")
        content = getattr(block, "content", "") or ""
        image_data = getattr(block, "image", None)
        image_path = None
        image_obj = None
        if image_data:
            image_path = image_data.get("path")
            image_obj = image_data.get("img")
        if label in {"image", "chart", "seal"} and image_path:
            content = image_path
        elif label not in {"table", "formula"}:
            content = _clean_text_content(content)

        blocks.append(
            {
                "type": label,
                "content": content,
                "normalized_content": _normalize_compare_text(content) if label not in {"image", "chart", "seal", "table", "formula"} else "",
                "signature_content": _normalize_signature_text(content) if label not in {"image", "chart", "seal", "table", "formula"} else "",
                "config": styles.get(label, {"size": 12, "indent": True}),
                "bbox": _normalize_bbox(getattr(block, "bbox", None)),
                "page_index": page_index,
                "order": order,
                "image_path": image_path,
                "image_obj": image_obj,
            }
        )

    return blocks


def _apply_paragraph_style(paragraph, config):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    paragraph.alignment = config.get("align", WD_ALIGN_PARAGRAPH.LEFT)
    if config.get("indent", False):
        paragraph.paragraph_format.first_line_indent = Inches(0.3)
    for run in paragraph.runs:
        _apply_run_style(run, config)


def _apply_run_style(run, config):
    from docx.oxml.ns import qn
    from docx.shared import Pt

    run.font.name = config.get("font", "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    run.font.size = Pt(config.get("size", 12))
    run.bold = config.get("bold", False)


def _add_picture(doc, block: dict, page_width: int):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches

    image_obj = block.get("image_obj")
    if image_obj is None:
        return False

    image = image_obj.convert("RGB") if getattr(image_obj, "mode", "RGB") != "RGB" else image_obj
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    buffer.seek(0)

    bbox = block.get("bbox") or [0, 0, image.width, image.height]
    bbox_width = max(int(bbox[2]) - int(bbox[0]), 1)
    width_ratio = bbox_width / max(float(page_width), 1.0)
    max_width = 5.4 if block.get("type") == "formula" else 6.2
    min_width = 0.9 if block.get("type") == "formula" else 1.5
    width_inches = min(max_width, max(min_width, width_ratio * 6.4))

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(buffer, width=Inches(width_inches))
    return True


def _normalize_formula_text(content: str) -> str:
    text = (content or "").strip()
    if text.startswith("$$") and text.endswith("$$"):
        text = text[2:-2].strip()
    if text.startswith("$") and text.endswith("$"):
        text = text[1:-1].strip()
    return text


_INLINE_MATH_PATTERN = re.compile(
    r"(\$[^$\n]+\$|=\s*\d+\s*/\s*\d+|\d+\s*/\s*\d+|P\([^)]+\)|[A-Za-z]\d{2,}|[A-Za-z]\s*\^\s*[0-9A-Za-z]+|[A-Za-z]\s*_\s*[0-9A-Za-z]+|\u221A\s*[0-9A-Za-z().+\-]+)"
)


def _to_latex_expression(token: str) -> str | None:
    raw = token or ""
    stripped = raw.strip()
    if stripped.startswith("$") and stripped.endswith("$") and len(stripped) >= 2:
        return _normalize_formula_text(stripped)

    compact = re.sub(r"\s+", "", stripped)
    if not compact:
        return None

    fraction_match = re.fullmatch(r"(\d+)/(\d+)", compact)
    if fraction_match:
        numerator, denominator = fraction_match.groups()
        return rf"\frac{{{numerator}}}{{{denominator}}}"

    if compact.startswith("P(") and compact.endswith(")"):
        return compact

    if compact.startswith("="):
        compact = compact[1:]
        fraction_match = re.fullmatch(r"(\d+)/(\d+)", compact)
        if fraction_match:
            numerator, denominator = fraction_match.groups()
            return rf"\frac{{{numerator}}}{{{denominator}}}"

    if compact.startswith("\u221A"):
        radicand = compact[1:]
        if radicand:
            return rf"\sqrt{{{radicand}}}"

    sup_match = re.fullmatch(r"([A-Za-z])\^([0-9A-Za-z]+)", compact)
    if sup_match:
        base, exponent = sup_match.groups()
        return rf"{base}^{{{exponent}}}"

    sub_match = re.fullmatch(r"([A-Za-z])_([0-9A-Za-z]+)", compact)
    if sub_match:
        base, subscript = sub_match.groups()
        return rf"{base}_{{{subscript}}}"

    digit_sub_match = re.fullmatch(r"([A-Za-z])(\d{2,})", compact)
    if digit_sub_match:
        base, subscript = digit_sub_match.groups()
        return rf"{base}_{{{subscript}}}"

    return None


def _render_inline_math(paragraph, content: str, config: dict) -> bool:
    content = unescape((content or "").strip())
    if not content:
        return False

    matches = list(_INLINE_MATH_PATTERN.finditer(content))
    if not matches:
        run = paragraph.add_run(content)
        _apply_run_style(run, config)
        return False

    cursor = 0
    inserted_math = False
    for match in matches:
        start, end = match.span()
        if start > cursor:
            run = paragraph.add_run(content[cursor:start])
            _apply_run_style(run, config)

        token = match.group(0)
        latex = _to_latex_expression(token)
        omml = _build_omml_from_latex(latex) if latex else None
        if omml is not None:
            paragraph._element.append(omml)
            inserted_math = True
        else:
            run = paragraph.add_run(token)
            _apply_run_style(run, config)
        cursor = end

    if cursor < len(content):
        run = paragraph.add_run(content[cursor:])
        _apply_run_style(run, config)

    return inserted_math


def _build_omml_from_latex(latex: str):
    try:
        from latex2mathml.converter import convert as latex_to_mathml
        from lxml import etree
    except Exception:
        return None

    xsl_path = _find_mml2omml_xsl()
    if xsl_path is None:
        return None

    try:
        mathml = latex_to_mathml(latex)
        mathml_tree = etree.fromstring(mathml.encode("utf-8"))
        transform = etree.XSLT(etree.parse(str(xsl_path)))
        omml_tree = transform(mathml_tree)
        return omml_tree.getroot()
    except Exception:
        return None


def _add_formula_block(doc, block: dict):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn

    latex = _normalize_formula_text(block.get("content") or "")
    if not latex:
        return False

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    omml = _build_omml_from_latex(latex)
    if omml is not None:
        paragraph._element.append(omml)
        return True

    run = paragraph.add_run(latex)
    run.font.name = "Cambria Math"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Cambria Math")
    try:
        run.font.math = True
    except Exception:
        pass
    return True


def _add_table(doc, html: str):
    rows = _parse_html_table(html)
    if not rows:
        return False

    max_cols = max(len(row) for row in rows)
    table = doc.add_table(rows=0, cols=max_cols)
    table.style = "Table Grid"
    cell_config = {"size": 10}
    previous_row_signature: tuple[str, ...] | None = None
    for row_cells in rows:
        cleaned_cells = [_clean_text_content(cell) for cell in row_cells]
        row_signature = tuple(_normalize_compare_text(cell) for cell in cleaned_cells)
        if previous_row_signature is not None and row_signature == previous_row_signature and any(row_signature):
            continue
        previous_row_signature = row_signature

        row = table.add_row().cells
        for idx in range(max_cols):
            cell_text = cleaned_cells[idx].strip() if idx < len(cleaned_cells) else ""
            paragraph = row[idx].paragraphs[0]
            _render_inline_math(paragraph, cell_text, cell_config)
            _apply_paragraph_style(paragraph, cell_config)
    return True


def _add_text_block(doc, block: dict):
    content = _clean_text_content(block.get("content") or "")
    if not content:
        return False

    config = block.get("config") or {}
    level = config.get("level")
    if isinstance(level, int) and 0 <= level <= 9:
        paragraph = doc.add_heading("", level=level)
    else:
        paragraph = doc.add_paragraph()
    _render_inline_math(paragraph, content, config)
    _apply_paragraph_style(paragraph, config)
    return True


def _should_skip_emitting_text(block: dict, recent_signatures: list[tuple[str, str]]):
    norm = block.get("signature_content") or block.get("normalized_content", "")
    label = block.get("type", "")
    if not norm or label in {"table", "formula", "image", "chart", "seal"}:
        return False

    return any(previous_norm == norm for _, previous_norm in recent_signatures)


def _should_insert_soft_break(blocks: list[dict]) -> bool:
    for block in blocks:
        label = block.get("type")
        content = (block.get("content") or "").strip()
        if not content and label not in {"image", "chart", "seal", "table", "formula"}:
            continue
        return label not in {"text", "vision_footnote"}
    return False


def _build_structure_engine(device: str, enable_formula: bool, enable_chart_recognition: bool = False):
    from paddleocr import PPStructureV3

    kwargs = {
        "device": device,
        "layout_detection_model_name": "PP-DocLayout_plus-L",
        "text_detection_model_name": "PP-OCRv5_server_det",
        "text_recognition_model_name": "PP-OCRv5_server_rec",
        "table_classification_model_name": "PP-LCNet_x1_0_table_cls",
        "wired_table_structure_recognition_model_name": "SLANeXt_wired",
        "wireless_table_structure_recognition_model_name": "SLANet_plus",
        "wired_table_cells_detection_model_name": "RT-DETR-L_wired_table_cell_det",
        "wireless_table_cells_detection_model_name": "RT-DETR-L_wireless_table_cell_det",
        "use_doc_orientation_classify": False,
        "use_doc_unwarping": False,
        "use_textline_orientation": False,
        "use_table_recognition": True,
        "use_formula_recognition": enable_formula,
        "use_chart_recognition": enable_chart_recognition,
        "use_seal_recognition": False,
        "use_region_detection": False,
    }
    if enable_formula:
        kwargs["formula_recognition_model_name"] = "PP-FormulaNet_plus-L"
    if enable_chart_recognition:
        kwargs["chart_recognition_model_name"] = "PP-Chart2Table"

    return PPStructureV3(**kwargs)


def run_paddle_structure(
    input_path: str,
    output_dir: str,
    requested_device: str | None = None,
    disable_formula: bool = False,
):
    try:
        from docx import Document
    except ImportError as exc:
        return {"success": False, "error": f"Failed to import Paddle dependencies: {exc}"}

    filename = os.path.basename(input_path)
    name, _ = os.path.splitext(filename)
    output_filepath = os.path.join(output_dir, f"{name}.docx")
    os.makedirs(output_dir, exist_ok=True)

    try:
        device = _resolve_device(requested_device)
        formula_enabled = _has_cached_model("PP-FormulaNet_plus-L") and not disable_formula
        engine = _build_structure_engine(device, formula_enabled, enable_chart_recognition=False)
        results = engine.predict(
            input_path,
            use_doc_orientation_classify=False,
            use_doc_unwarping=False,
            use_textline_orientation=False,
            use_table_recognition=True,
            use_formula_recognition=formula_enabled,
            use_chart_recognition=False,
            use_seal_recognition=False,
            use_region_detection=False,
            use_table_orientation_classify=False,
        )

        merged_doc = Document()
        first_page = True
        for page_result in results:
            blocks = _deduplicate_page_blocks(_extract_page_blocks(page_result))
            if not blocks:
                continue
            recent_text_signatures: list[tuple[str, str]] = []

            if not first_page and _should_insert_soft_break(blocks):
                merged_doc.add_paragraph()

            page_width = int(page_result.get("width") or 1200)
            for block in blocks:
                label = block.get("type")
                content = (block.get("content") or "").strip()

                if label == "table" and content:
                    if _add_table(merged_doc, content):
                        continue

                if label == "formula":
                    if _add_formula_block(merged_doc, block):
                        continue

                if label in {"image", "chart", "seal"} and block.get("image_obj") is not None:
                    if _add_picture(merged_doc, block, page_width):
                        continue

                if _should_skip_emitting_text(block, recent_text_signatures):
                    continue

                if _add_text_block(merged_doc, block):
                    recent_text_signatures.append(
                        (
                            block.get("type", ""),
                            block.get("signature_content") or block.get("normalized_content", ""),
                        )
                    )
                    recent_text_signatures = recent_text_signatures[-20:]

            first_page = False

        if first_page:
            return {
                "success": False,
                "error": "PPStructureV3 returned no structured blocks.",
                "device": device,
                "formula_enabled": formula_enabled,
            }

        merged_doc.save(output_filepath)
        return {
            "success": True,
            "output_path": output_filepath,
            "converter_used": "paddle_structure_v4",
            "device": device,
            "formula_enabled": formula_enabled,
        }
    except Exception as exc:
        return {
            "success": False,
            "error": f"Exception during PPStructureV3 processing: {exc}",
            "device": requested_device or os.environ.get("PADDLE_STRUCTURE_DEVICE") or "auto",
            "traceback": traceback.format_exc(),
        }


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, help="Input PDF path")
    parser.add_argument("--output_dir", required=True, help="Output directory path")
    parser.add_argument("--device", choices=["auto", "cpu", "gpu:0"], default="auto", help="Execution device")
    parser.add_argument("--disable-formula", action="store_true", help="Disable formula recognition for stability fallback")
    args = parser.parse_args()

    result = run_paddle_structure(
        args.input,
        args.output_dir,
        requested_device=args.device,
        disable_formula=args.disable_formula,
    )
    print(json.dumps(result, ensure_ascii=False))
