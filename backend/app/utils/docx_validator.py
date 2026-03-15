import xml.etree.ElementTree as ET
import re
from docx import Document
from ..core.logger import logger

from typing import Dict, Any

def validate_docx_quality(filepath: str) -> Dict[str, Any]:
    """
    Validates the quality of a DOCX file generated from PDF, generating a detailed report.
    Checks for pseudo-success structures and common layout recovery issues.
    """
    try:
        doc = Document(filepath)
    except Exception as e:
        logger.error(f"Error reading docx for validation: {e}")
        return {"final_quality_level": "poor", "warnings": ["Failed to parse generated DOCX file."]}

    paragraphs = list(doc.paragraphs)
    text_paras = [p.text.strip() for p in paragraphs if p.text.strip()]
    num_text_paras = len(text_paras)

    # Element counts
    drawing_count = 0
    shape_count = 0
    textbox_count = 0
    behind_doc_count = 0

    # Heuristic layout metrics
    short_para_count = 0
    math_layout_score = 1.0
    spacing_score = 1.0
    heading_structure_score = 0.0

    headings = 0
    math_anomalies = 0
    spacing_anomalies = 0

    for p in paragraphs:
        p_xml = p._element.xml
        drawing_count += p_xml.count('<w:drawing>')
        shape_count += p_xml.count('<v:shape')
        textbox_count += p_xml.count('<v:textbox') + p_xml.count('<w:txbxContent')
        behind_doc_count += p_xml.count('behindDoc="1"')

        text = p.text.strip()
        if text:
            # Check for very short isolated paragraphs (often broken lines)
            if len(text) < 15 and not any(text.startswith(c) for c in ('A', 'B', 'C', 'D', '(', '例', '变')):
                short_para_count += 1
            # Check for math fragmentation (e.g. single P, (, A, ) scattered)
            if text in ['P', '(', ')', '/'] or re.match(r'^\d+\s*/\s*\d+$', text):
                math_anomalies += 1
            # Check for spacing anomalies (e.g. English words separated by spaces like E x a m p l e)
            if re.search(r'\b[A-Za-z]\s+[A-Za-z]\b', text):
                spacing_anomalies += 1
            # Check for heading structures (【】, 例, 变式)
            if re.match(r'^(【.*?】|专题|例\s*\d+|变式|考点)', text):
                headings += 1

    table_count = len(doc.tables)

    # Calculate scores (0.0 to 1.0)
    if num_text_paras > 0:
        math_layout_score = max(0.0, 1.0 - (math_anomalies / (num_text_paras * 0.5)))
        spacing_score = max(0.0, 1.0 - (spacing_anomalies / (num_text_paras * 0.5)))
        heading_structure_score = min(1.0, headings / 5.0) # Assume 5+ headings is good structure

    logger.info(
        f"DOCX Validation: {num_text_paras} standard paragraphs, {table_count} tables, "
        f"{drawing_count} drawings, {shape_count} shapes, "
        f"{textbox_count} textboxes, {behind_doc_count} behindDoc."
    )

    # Construct report
    report = {
        "visible_text_ok": num_text_paras > 5,
        "paragraph_recovery_score": 1.0 - (short_para_count / num_text_paras) if num_text_paras > 0 else 0,
        "table_recovery_score": table_count,
        "math_layout_score": math_layout_score,
        "spacing_score": spacing_score,
        "heading_structure_score": heading_structure_score,
        "warnings": []
    }

    # Evaluate final quality
    # Pseudo-success (blank)
    if not report["visible_text_ok"] and (drawing_count > 10 or shape_count > 10 or textbox_count > 10):
        report["final_quality_level"] = "failed"
        report["warnings"].append("转换失败：生成的 Word 结构不兼容或不可见，已拦截此结果。")
    elif num_text_paras < 3 and not table_count:
        report["final_quality_level"] = "poor"
        report["warnings"].append("提取出的有效文本极少，可能是一个图片型 PDF。")
    elif report["paragraph_recovery_score"] < 0.4:
        report["final_quality_level"] = "acceptable"
        report["warnings"].append("段落结构恢复较差，存在较多换行断裂。")
    elif report["math_layout_score"] < 0.6:
        report["final_quality_level"] = "acceptable"
        report["warnings"].append("部分数学公式、分数或区间可能存在排版粘连或断裂。")
    elif report["spacing_score"] < 0.7:
        report["final_quality_level"] = "acceptable"
        report["warnings"].append("中英文混排、特殊符号间距存在一定异常。")
    elif report["table_recovery_score"] == 0 and num_text_paras > 50:
         # Large document, no tables - might be ok, but note it
         report["final_quality_level"] = "good"
         if "表" in " ".join(text_paras):
             # Just a heuristic to guess if there should be tables
             pass
    else:
        report["final_quality_level"] = "good"

    return report
